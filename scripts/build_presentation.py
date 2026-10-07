"""build_presentation.py -- presentation/InfraDriftGuard_Phase1_Review.pptx"""

import json
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import report_content as RC  # noqa: E402

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
ORANGE = RGBColor(0xED, 0x71, 0x00)
GREY = RGBColor(0x55, 0x5F, 0x6D)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def bar(slide):
    shp = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.12))
    shp.fill.solid()
    shp.fill.fore_color.rgb = ORANGE
    shp.line.fill.background()


def title(slide, text, sub=None):
    bar(slide)
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12), Inches(0.9)).text_frame
    tb.text = text
    tb.paragraphs[0].runs[0].font.size = Pt(30)
    tb.paragraphs[0].runs[0].font.bold = True
    tb.paragraphs[0].runs[0].font.color.rgb = NAVY
    if sub:
        p = tb.add_paragraph()
        p.text = sub
        p.runs[0].font.size = Pt(15)
        p.runs[0].font.color.rgb = GREY


def bullets(slide, items, x=0.7, y=1.6, w=12, h=5.4, size=19):
    tf = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(it, tuple):
            r = p.add_run()
            r.text = it[0] + "  "
            r.font.bold = True
            r.font.color.rgb = NAVY
            r.font.size = Pt(size)
            r2 = p.add_run()
            r2.text = it[1]
            r2.font.size = Pt(size - 2)
        else:
            p.text = "•  " + it
            p.runs[0].font.size = Pt(size)
        p.space_after = Pt(10)


def image(slide, path, x, y, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    slide.shapes.add_picture(path, Inches(x), Inches(y), **kw)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# 1 title
s = prs.slides.add_slide(BLANK)
bar(s)
tf = s.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.7), Inches(4)).text_frame
tf.word_wrap = True
tf.text = "InfraDriftGuard"
tf.paragraphs[0].runs[0].font.size = Pt(54)
tf.paragraphs[0].runs[0].font.bold = True
tf.paragraphs[0].runs[0].font.color.rgb = NAVY
for txt, size, col in ((RC.TEAM["project_title"], 20, GREY), ("", 10, GREY),
                       (RC.TEAM["course"] + " · Project Phase-I · " + RC.TEAM["instructor"], 16, GREY),
                       ("  ·  ".join(f"{st['name']} ({st['reg']})" for st in RC.TEAM["students"]), 16, NAVY)):
    p = tf.add_paragraph()
    p.text = txt
    if p.runs:
        p.runs[0].font.size = Pt(size)
        p.runs[0].font.color.rgb = col

# 2 problem
s = prs.slides.add_slide(BLANK)
title(s, "The problem: configuration drift")
bullets(s, [
    ("Declared ≠ deployed.", "Terraform describes the infrastructure, but console edits, hot-fixes and scripts change "
     "the live AWS account afterwards."),
    ("Drift causes breaches.", "A public S3 bucket, SSH open to 0.0.0.0/0, or an IAM role escalated to Action:* can "
     "each expose the whole account."),
    ("Today's tools fall short.", "Static scanners never see the live account; terraform plan / AWS Config say "
     "“something changed” without saying how dangerous; ML/LLM tools are opaque or hallucinate."),
    ("Goal.", "Detect drift continuously, rate its risk, explain why, and respond within the free tier."),
])

# 3 literature
s = prs.slides.add_slide(BLANK)
title(s, "Literature survey — 15 papers (2023–2026)", "IEEE · ACM · Springer · Frontiers · Scopus")
bullets(s, [
    ("Static IaC analysis", "GLITCH (ASE'23), GASEL (MSR'23), SLI-KUBE (TOSEM'23), TerraMetrics (ICPC'24), "
     "Verdet (EMSE'25), SlsDetector (TOSEM'26) → miss post-deployment changes"),
    ("Drift & state", "Thiyagarajan (IJCNIS'24), Hassan (FSE'24), RIVA (ACM'26) → no severity, or costly LLM agents"),
    ("Repair", "InfraFix (ISSTA'25), LLM repair (SecDev'24) → fixes code; 20% hallucinated fixes"),
    ("ML / XAI for security", "Begoug (MSR'24), Saleh (CCSW'24), Rjoub (TNSM'23), Mohale (Frontiers'25) → "
     "explanations needed, not yet applied to cloud drift"),
], size=18)

# 4 gap + objectives
s = prs.slides.add_slide(BLANK)
title(s, "Research gap → objectives")
bullets(s, RC.COMBINED_GAP[:5], x=0.6, w=6.0, size=15)
bullets(s, [f"{i}. {t}" for i, (t, _) in enumerate(RC.OBJECTIVES, 1)], x=7.0, w=5.8, size=17)

# 5 architecture
s = prs.slides.add_slide(BLANK)
title(s, "AWS cloud architecture")
image(s, os.path.join(ROOT, "architecture", "AWS_Architecture.png"), 0.9, 1.25, h=6.1)
notes(s, "Walk the two paths: A-G detection path triggered hourly by EventBridge; 1-5 user path through Cognito and "
         "API Gateway. Terraform renders the desired state into the Lambda environment.")

# 6 system
s = prs.slides.add_slide(BLANK)
title(s, "Complete system architecture")
image(s, os.path.join(ROOT, "architecture", "System_Architecture.png"), 1.3, 1.2, h=6.2)

# 7 novelty
s = prs.slides.add_slide(BLANK)
title(s, "Novelty")
bullets(s, [
    ("Runtime, not just code:", "compares live AWS with Terraform's desired state every hour"),
    ("Risk level per drift:", "9 security features → Low / Medium / High / Critical"),
    ("Explainable:", "exact SHAP explanations for every prediction"),
    ("No guessing:", "never-seen changes are flagged Review, not forced into a class"),
    ("Fully serverless AWS:", "EventBridge → Lambda → DynamoDB / SNS / CloudWatch, Cognito-secured API"),
    ("Safe automation:", "opt-in revert of public S3 access and internet-open SSH/RDP only"),
    ("Free-tier by design:", "210 KB dependency-free model, $1 budget guard-rail, least-privilege IAM"),
], size=19)

# 8 dataset + model
s = prs.slides.add_slide(BLANK)
title(s, "Dataset & model")
bullets(s, [
    ("711 records,", "20 drift scenarios (7 S3, 7 security group, 6 IAM), balanced across 4 risk levels"),
    ("9 features:", "resource type, changed attribute, sensitivity, public exposure, encryption change, privilege "
     "change, port exposure, magnitude, drift frequency"),
    ("Random Forest (200 trees)", "+ Decision Tree baseline; SHAP TreeExplainer"),
    ("Exported to 210 KB JSON", "→ pure-Python inference in Lambda, identical to scikit-learn on all 711 records"),
], w=6.6, size=17)
image(s, os.path.join(ROOT, "results", "shap", "global_bar.png"), 7.5, 1.4, w=5.4)

# 9 results
m = json.load(open(os.path.join(ROOT, "results", "metrics.json")))
a, b = m["split_A_stratified_random"], m["split_B_scenario_group_aware"]
s = prs.slides.add_slide(BLANK)
title(s, "Results")
rows = [("Evaluation", "Accuracy", "Macro F1"),
        ("RF – record-level split", f"{a['random_forest']['accuracy']:.3f}", f"{a['random_forest']['f1_macro']:.3f}"),
        ("RF – scenarios held out", f"{b['random_forest']['accuracy']:.3f}", f"{b['random_forest']['f1_macro']:.3f}"),
        ("DT – scenarios held out", f"{b['decision_tree']['accuracy']:.3f}", f"{b['decision_tree']['f1_macro']:.3f}")]
tbl = s.shapes.add_table(4, 3, Inches(0.7), Inches(1.6), Inches(6.4), Inches(2.0)).table
for r, row in enumerate(rows):
    for c, v in enumerate(row):
        tbl.cell(r, c).text = v
        tbl.cell(r, c).text_frame.paragraphs[0].runs[0].font.size = Pt(15)
bullets(s, [
    "218 automated tests pass (moto-simulated AWS)",
    "Public SSH → Critical + email; auto-remediation verified",
    "Unseen change → Review (never guessed)",
    "Honest limitation: generalisation to new scenarios",
], y=4.0, w=6.6, size=16)
image(s, os.path.join(ROOT, "results", "shap", "local_critical.png"), 7.4, 1.4, w=5.6)

# 10 demo + future
s = prs.slides.add_slide(BLANK)
title(s, "Live demo · GitHub · next steps")
bullets(s, [
    ("Demo:", "terraform apply → open SSH 0.0.0.0/0 in console → Scan now → Critical email → auto-revert"),
    ("GitHub:", f"{RC.TEAM['repo']} — main ← develop ← feature/student1, feature/student2, tag v1.0-Phase1"),
    ("Cost:", "$0 within free tier; $1 AWS Budget guard-rail; terraform destroy when done"),
    ("Phase II:", "real drift data, CloudTrail event-driven scans, more services (RDS, KMS, VPC), multi-account"),
], size=18)

out = os.path.join(ROOT, "presentation", "InfraDriftGuard_Phase1_Review.pptx")
os.makedirs(os.path.dirname(out), exist_ok=True)
prs.save(out)
print("Wrote", out)
