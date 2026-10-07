"""
build_docs.py -- generates every Word document in docs/ (plus
dataset/Dataset_Details.docx) from scripts/report_content.py, the saved
results/ and the architecture/ diagrams.

    python scripts/build_docs.py
"""

import json
import os
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import report_content as RC  # noqa: E402

DOCS = os.path.join(ROOT, "docs")
ARCH = os.path.join(ROOT, "architecture")
RES = os.path.join(ROOT, "results")
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------

def new_doc():
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.2))
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.15
    for lvl, size in ((1, 16), (2, 13), (3, 11.5)):
        hs = d.styles[f"Heading {lvl}"]
        hs.font.name = "Calibri"
        hs.font.size = Pt(size)
        hs.font.color.rgb = ACCENT
        hs.font.bold = True
        hs.element.rPr.rFonts.set(qn("w:asciiTheme"), "")
    _footer_page_numbers(sec)
    return d


def _footer_page_numbers(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    run.font.size = Pt(9)


def para(d, text, bold=False, italic=False, size=None, align=None, color=None, after=None):
    p = d.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    if align:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def rich(d, parts, style=None):
    """parts: list of (text, bold)"""
    p = d.add_paragraph(style=style)
    for text, bold in parts:
        p.add_run(text).bold = bold
    return p


def bullets(d, items, numbered=False):
    for it in items:
        if isinstance(it, tuple):
            rich(d, [(it[0] + ": ", True), (it[1], False)], style="List Number" if numbered else "List Bullet")
        else:
            d.add_paragraph(it, style="List Number" if numbered else "List Bullet")


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def table(d, header, rows, widths_cm, font=9.5):
    t = d.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(font)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, "1F3A5F")
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(font)
            if ri % 2:
                shade(cells[i], "EEF3F8")
    for row in t.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)
    # repeat header row
    hdr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    hdr.append(th)
    d.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figure(d, path, caption, width_cm=16.5):
    if not os.path.exists(path):
        para(d, f"[missing figure: {os.path.relpath(path, ROOT)}]", italic=True)
        return
    d.add_picture(path, width=Cm(width_cm))
    d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(d, caption, italic=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0x55, 0x5F, 0x6D))


def page_break(d):
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def landscape_section(d):
    sec = d.add_section()
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21)
    sec.left_margin = sec.right_margin = Cm(1.6)
    return sec


def portrait_section(d):
    sec = d.add_section()
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    return sec


def doc_header(d, title, subtitle=None):
    para(d, RC.TEAM["course"] + "  ·  Project Phase-I", size=10, color=RGBColor(0x55, 0x5F, 0x6D), after=0)
    para(d, RC.TEAM["short_name"], bold=True, size=11, color=ACCENT, after=2)
    d.add_heading(title, level=1)
    if subtitle:
        para(d, subtitle, italic=True, size=10)


def students_line():
    return "; ".join(f"{s['name']} ({s['reg']})" for s in RC.TEAM["students"])


# ---------------------------------------------------------------------------
# reusable sections
# ---------------------------------------------------------------------------

def sec_abstract(d, h=1):
    d.add_heading("Abstract", level=h)
    para(d, RC.ABSTRACT, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    rich(d, [("Keywords: ", True), ("Infrastructure as Code, configuration drift, AWS, Terraform, serverless, "
                                    "Random Forest, SHAP, explainable AI, cloud security", False)])


def sec_literature(d, h=1, landscape=True):
    if landscape:
        landscape_section(d)
    d.add_heading("Literature Survey", level=h)
    para(d, "Fifteen peer-reviewed papers published 2023–2026 (IEEE, ACM, Springer, Frontiers and Scopus-indexed "
            "journals) were studied. Papers 1–8 were studied by Student 1 and papers 9–15 by Student 2, following "
            "the team-of-two allocation.")
    rows = []
    for i, p in enumerate(RC.PAPERS, 1):
        rows.append([f"[{i}] {p['short']}", p["method"], p["dataset"], p["advantages"], p["limitations"], p["gap"]])
    table(d, ["Paper", "Method", "Dataset", "Advantages", "Limitations", "Research Gap"], rows,
          [3.6, 5.0, 3.6, 4.6, 4.6, 4.6], font=8.5)
    d.add_heading("Summary of the survey", level=h + 1)
    para(d, RC.SURVEY_SUMMARY, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    if landscape:
        portrait_section(d)


def sec_research_gap(d, student, h=1):
    lo, hi = student["papers"]
    d.add_heading(f"Research Gap Analysis — {student['id']} (Papers {lo}–{hi})", level=h)
    para(d, f"Prepared by {student['name']} ({student['reg']}). Gaps are stated in our own words based on our "
            "reading of each paper, not copied from the papers.", italic=True, size=10)
    for i in range(lo, hi + 1):
        p = RC.PAPERS[i - 1]
        d.add_heading(f"[{i}] {p['short']}", level=h + 1)
        for label, key in (("Existing method", "method"), ("Advantages", "advantages"),
                           ("Limitations", "limitations"), ("Research gap", "gap"),
                           ("Possible improvement (addressed in InfraDriftGuard)", "improvement")):
            rich(d, [(label + ": ", True), (p[key], False)])


def sec_combined_gap(d, h=1):
    d.add_heading("Consolidated Research Gap", level=h)
    para(d, "Taken together, the surveyed work leaves the following gaps, which define the scope of this project:")
    bullets(d, RC.COMBINED_GAP)


def sec_objectives(d, h=1):
    d.add_heading("Project Objectives", level=h)
    bullets(d, RC.OBJECTIVES, numbered=True)


def sec_novelty(d, h=1):
    d.add_heading("Novelty Summary", level=h)
    para(d, "What makes InfraDriftGuard different from existing work:", bold=True)
    for title, text in RC.NOVELTY:
        rich(d, [(title + ". ", True), (text, False)], style="List Bullet")


def sec_architecture(d, h=1):
    d.add_heading("Proposed Architecture", level=h)
    d.add_heading("Diagram 1 — AWS Cloud Architecture", level=h + 1)
    figure(d, os.path.join(ARCH, "AWS_Architecture.png"), "Figure 1. AWS cloud architecture of InfraDriftGuard.")
    rows = [
        ("Data flow", "EventBridge → Detector Lambda → (boto3 read-only) S3 / EC2 SG / IAM → Detector → DynamoDB, "
                      "CloudWatch, SNS. Dashboard → Cognito → API Gateway → API Lambda → DynamoDB."),
        ("Storage", "DynamoDB table drift-findings (partition key resource_id, sort key detected_at), encrypted at "
                    "rest; model shipped inside the Lambda package."),
        ("Processing", "Detector Lambda: collect → normalise → detect_drift → extract 9 features → pure-Python "
                       "Random Forest → risk label; API Lambda for queries and on-demand scans."),
        ("Authentication", "Amazon Cognito user pool (admin-created users, strong password policy, 1-hour "
                           "tokens); API Gateway JWT authorizer on every route; IAM least-privilege roles."),
        ("Notifications", "Amazon SNS e-mail for High/Critical/Review drift and for CloudWatch alarms; AWS "
                          "Budgets e-mails for cost."),
        ("Monitoring", "CloudWatch Logs (7-day retention), custom metric InfraDriftGuard/DriftDetected by "
                       "RiskLevel, alarms on Lambda errors and Critical drift, CloudWatch dashboard."),
    ]
    table(d, ["Aspect", "How it is realised"], rows, [3.5, 13.0])
    page_break(d)
    d.add_heading("Diagram 2 — Complete System Architecture", level=h + 1)
    figure(d, os.path.join(ARCH, "System_Architecture.png"), "Figure 2. Complete system architecture: offline "
           "model development, online detection, response and presentation.")
    para(d, "The system has three layers. The offline layer builds the scenario library, generates and validates "
            "the dataset, trains the Random Forest and Decision Tree, produces SHAP explanations and exports the "
            "model to JSON. The online layer runs in AWS Lambda: Terraform supplies the desired state, collectors "
            "read live configuration, normalisers map it to a common schema, the drift engine diffs the two and "
            "the feature extractor and classifier rate the risk. The response and presentation layer stores "
            "findings, alerts, optionally remediates, publishes metrics and serves the results to users.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    page_break(d)
    d.add_heading("Scan workflow", level=h + 1)
    figure(d, os.path.join(ARCH, "Workflow.png"), "Figure 3. Decision flow of a single drift scan.", width_cm=14)


def sec_dataset(d, h=1):
    D = RC.DATASET
    d.add_heading("Dataset Details", level=h)
    rows = [("Dataset name", D["name"]), ("Source", D["source"]), ("URL", D["url"]), ("Size", D["size"]),
            ("Number of records", D["records"]), ("Number of features", D["features"]), ("Data type", D["types"]),
            ("License", D["license"]), ("Purpose of using the dataset", D["purpose"])]
    table(d, ["Field", "Details"], rows, [4.0, 12.5])
    d.add_heading("Features", level=h + 1)
    table(d, ["Feature", "Type", "Meaning"], [
        ("resource_type", "Categorical", "s3_bucket / security_group / iam_policy"),
        ("changed_attribute", "Categorical", "Which attribute(s) drifted, e.g. block_public_access"),
        ("security_sensitivity", "Ordinal", "Intrinsic sensitivity of the attribute (Low/Medium/High)"),
        ("public_exposure", "Boolean", "Did the change expose the resource to the internet?"),
        ("encryption_change", "Boolean", "Was encryption weakened or removed?"),
        ("privilege_change", "Boolean", "Were permissions escalated or a control removed?"),
        ("port_exposure", "Ordinal", "none / internal / 0.0.0.0/0"),
        ("change_magnitude", "Ordinal", "None/Low/Medium/High size of the change"),
        ("drift_frequency", "Integer", "How often this resource has drifted before (from DynamoDB history)"),
        ("risk_label (target)", "Ordinal", "Low / Medium / High / Critical"),
    ], [4.0, 2.5, 10.0])
    d.add_heading("Drift scenarios", level=h + 1)
    table(d, ["ID", "Resource", "Drift", "Risk"], D["scenarios"], [1.6, 3.2, 8.8, 2.9], font=9)
    d.add_heading("Preprocessing required", level=h + 1)
    bullets(d, D["preprocessing"])


def sec_aws(d, h=1):
    d.add_heading("AWS Services Planning", level=h)
    table(d, ["AWS Service", "Purpose in InfraDriftGuard", "Free-tier position"], RC.AWS_SERVICES, [4.0, 8.5, 4.0])
    para(d, "Note: AWS accounts created after 15 July 2025 use the credit-based Free Plan (up to US$200 in credits "
            "for six months) in addition to the always-free allowances listed above. The $1 AWS Budget alarm "
            "e-mails the team long before any charge could matter.", italic=True, size=9.5)


def _load_json(*parts):
    path = os.path.join(RES, *parts)
    with open(path) as f:
        return json.load(f)


def sec_implementation(d, h=1):
    d.add_heading("Implementation", level=h)
    d.add_heading("Repository modules", level=h + 1)
    table(d, ["Module", "Responsibility"], [
        ("src/backend/drift_engine.py", "Generic recursive diff of desired vs actual state"),
        ("src/backend/feature_extraction.py", "21-rule knowledge base converting each change to 9 features"),
        ("src/backend/scenario_definitions.py, dataset_generator.py", "20 scenarios → 711-record dataset"),
        ("src/ml_model/ml_pipeline.py", "Preprocessing, two splits, Random Forest + Decision Tree, metrics"),
        ("src/ml_model/shap_explainability.py", "SHAP TreeExplainer, global/local plots, additivity check"),
        ("src/ml_model/export_model.py", "Exports the Random Forest + encoder to dependency-free JSON"),
        ("src/aws/collectors/*", "boto3 collection and normalisation for S3, security groups, IAM roles"),
        ("src/aws/lambda/detector_handler.py", "Scheduled detector: diff, classify, store, alert, remediate"),
        ("src/aws/lambda/api_handler.py", "REST API behind API Gateway + Cognito"),
        ("src/aws/lambda/rf_predict.py", "Pure-Python Random Forest inference"),
        ("src/aws/terraform/*.tf", "All AWS infrastructure + rendered desired state"),
        ("src/frontend/app.py, cloud_client.py", "Streamlit dashboard: offline analyzer + live AWS findings"),
        ("tests/", "219 automated tests (pytest + moto simulated AWS)"),
    ], [6.5, 10.0], font=9)
    d.add_heading("Key design decisions", level=h + 1)
    bullets(d, [
        ("Terraform as single source of truth", "the desired state the Lambda compares against is rendered by "
         "Terraform from the same resources it creates, so code and monitor cannot disagree."),
        ("Dependency-free model", "200 decision trees are serialised (5,472 nodes, 210 KB) and evaluated in pure "
         "Python; a parity test confirms identical probabilities to scikit-learn for all 711 records."),
        ("Real drift_frequency", "computed from the resource's own finding history in DynamoDB."),
        ("Unknown changes are never guessed", "changes outside the knowledge base are labelled Review and "
         "escalated."),
        ("Failure isolation", "an error on one resource is logged and the scan continues with the others."),
        ("Least privilege", "collection permissions are read-only and scoped; remediation permissions exist only "
         "when auto_remediate = true."),
    ])


def sec_results(d, h=1):
    d.add_heading("Results", level=h)
    m = _load_json("metrics.json")
    a, b = m["split_A_stratified_random"], m["split_B_scenario_group_aware"]

    def row(split, name, key):
        r = split[key]
        return [name, f"{r['accuracy']:.3f}", f"{r['precision_macro']:.3f}", f"{r['recall_macro']:.3f}",
                f"{r['f1_macro']:.3f}"]

    d.add_heading("Classification performance", level=h + 1)
    table(d, ["Model / split", "Accuracy", "Precision (macro)", "Recall (macro)", "F1 (macro)"], [
        row(a, f"Random Forest – Split A (record-level, n_test={a['n_test']})", "random_forest"),
        row(a, "Decision Tree – Split A", "decision_tree"),
        row(b, f"Random Forest – Split B (scenario held out, n_test={b['n_test']})", "random_forest"),
        row(b, "Decision Tree – Split B", "decision_tree"),
    ], [7.0, 2.2, 2.6, 2.4, 2.3])
    para(d, "Interpretation. Split A meets the ≥95% macro-F1 objective, but every scenario appears in both train and "
            "test, so it largely measures how well the model reproduces the labelling rules. Split B holds out "
            "entire scenarios (E1, E2, S3id, S5) and is far harder: two (resource type, risk) combinations never "
            "appear in training. We report both numbers deliberately; improving generalisation to unseen "
            "scenarios (more scenario templates, real drift data) is the main item of future work.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    figure(d, os.path.join(RES, "confusion_matrix_split_A_stratified_random_random_forest.png"),
           "Figure 4. Confusion matrix, Random Forest, Split A.", width_cm=9.5)
    figure(d, os.path.join(RES, "confusion_matrix_split_B_scenario_group_aware_random_forest.png"),
           "Figure 5. Confusion matrix, Random Forest, Split B (scenario held out).", width_cm=9.5)

    g = _load_json("shap", "global_feature_importance.json")
    d.add_heading("Explainability (SHAP)", level=h + 1)
    top = g["mean_abs_shap_overall_across_all_classes"][:6]
    table(d, ["Rank", "Feature", "Mean |SHAP|"],
          [(i + 1, t["feature"].split("__", 1)[1], f"{t['mean_abs_shap']:.4f}") for i, t in enumerate(top)],
          [1.5, 9.0, 3.5])
    para(d, f"SHAP additivity holds: the largest difference between base value + SHAP contributions and the model's "
            f"predicted probability is {g['additivity_check']['max_abs_diff_from_predict_proba']:.1e}. The two most "
            "influential features are security_sensitivity and change_magnitude, matching the security reasoning "
            "the scenarios were designed around.")
    figure(d, os.path.join(RES, "shap", "global_summary.png"), "Figure 6. SHAP beeswarm for the Critical class.",
           width_cm=13)
    figure(d, os.path.join(RES, "shap", "local_critical.png"),
           "Figure 7. Local SHAP explanation of a Critical prediction (SSH opened to 0.0.0.0/0).", width_cm=13)

    d.add_heading("Cloud pipeline verification", level=h + 1)
    para(d, "The AWS pipeline was tested end-to-end against moto (an in-memory simulation of the AWS APIs), "
            "covering the exact resources Terraform deploys:")
    table(d, ["Test scenario", "Expected", "Result"], [
        ("Fresh deployment, no manual changes", "No drift on all 3 resources, no alert", "Pass"),
        ("SSH (22) opened to 0.0.0.0/0", "Critical, SNS alert, finding stored", "Pass"),
        ("S3 public-access block disabled + auto-remediate", "Critical, block re-enabled, next scan clean", "Pass"),
        ("Public SSH + auto-remediate", "Rule revoked, next scan clean", "Pass"),
        ("IAM inline policy escalated to Action:*/Resource:*", "Critical", "Pass"),
        ("MFA condition removed from trust policy", "High", "Pass"),
        ("Versioning enabled outside Terraform", "Low/Medium, no alert", "Pass"),
        ("Same drift seen in 3 scans", "drift_frequency = 0, 1, 2", "Pass"),
        ("HTTPS rule deleted (attribute never modelled)", "Review + alert (not guessed)", "Pass"),
        ("Monitored bucket missing", "Error recorded, other resources still scanned", "Pass"),
        ("API: list findings / history / unknown resource", "200 / 200 / 404", "Pass"),
        ("Exported model vs scikit-learn on 711 records", "Identical probabilities", "Pass"),
    ], [7.5, 6.0, 2.0], font=9)
    para(d, "Full suite: 218 tests passed, 1 skipped (the skipped test needs a real AWS account and runs with "
            "RUN_REAL_AWS_TESTS=1).", bold=True)
    for name, cap in (("dashboard_input_S1.jpg", "Figure 8. Streamlit dashboard – scenario S1 input."),
                      ("dashboard_prediction_critical.jpg", "Figure 9. Dashboard prediction: Critical.")):
        figure(d, os.path.join(RES, "screenshots", name), cap, width_cm=14)


def sec_deploy(d, h=1):
    d.add_heading("Deployment and Cost Control", level=h)
    bullets(d, [
        "Create an AWS account, enable MFA on the root user and create an IAM admin user for daily use.",
        "Install Terraform ≥ 1.5 and the AWS CLI; run aws configure.",
        "cd src/aws/terraform && terraform init && terraform apply -var alert_email=<you@example.com>",
        "Confirm the SNS subscription e-mail; sign in to the dashboard with the temporary Cognito password.",
        "terraform output -json > outputs.json; streamlit run src/frontend/app.py → Live AWS findings tab.",
        "Demo drift: open SSH to 0.0.0.0/0 in the console, click Scan now, receive the Critical alert e-mail.",
        "When finished: terraform destroy (removes everything, cost returns to $0).",
    ], numbered=True)
    para(d, "Expected monthly cost for the demo: US$0.00 within free-tier limits (~720 Lambda invocations, a few "
            "hundred DynamoDB writes, a handful of e-mails). The AWS Budget alarm triggers at $0.50 actual or "
            "$1 forecast.", italic=True)


def sec_github(d, h=1):
    d.add_heading("GitHub Repository and Contributions", level=h)
    rich(d, [("Repository: ", True), (RC.TEAM["repo"], False)])
    para(d, "Branches: main ← develop ← feature/student1, feature/student2. Each student works only on their "
            "feature branch, opens pull requests into develop, reviews the partner's PRs, and develop is merged "
            "into main and tagged v1.0-Phase1 once stable.")
    s1, s2 = RC.TEAM["students"]
    table(d, ["Activity", s1["id"], s2["id"]], RC.CONTRIBUTIONS, [8.5, 4.0, 4.0])


def sec_conclusion(d, h=1):
    d.add_heading("Conclusion and Future Work", level=h)
    para(d, "InfraDriftGuard closes the gap between static IaC analysis and runtime cloud security: it detects "
            "drift between Terraform and live AWS resources every hour, rates each drift with an explainable "
            "Random Forest, never guesses on unfamiliar changes, alerts within one scan cycle and can safely "
            "revert the most dangerous changes – all in a serverless design that costs nothing within the free "
            "tier.", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    para(d, "Future work (Phase II):", bold=True)
    bullets(d, [
        "Collect real, anonymised drift events (e.g. CloudTrail-triggered scans) to train beyond synthetic data and "
        "improve the scenario-held-out score.",
        "Event-driven detection via CloudTrail + EventBridge rules for near-real-time response.",
        "Extend coverage to RDS, KMS keys, Lambda configuration and VPC network ACLs.",
        "Multi-account monitoring through AWS Organizations and cross-account read-only roles.",
        "Generate pull requests that update Terraform when a drift is intentional (accept vs revert workflow).",
    ])


def sec_references(d, h=1):
    d.add_heading("References", level=h)
    for i, p in enumerate(RC.PAPERS, 1):
        para(d, f"[{i}] {p['ref']}", size=10)


def title_page(d, title):
    for _ in range(5):
        d.add_paragraph()
    para(d, RC.TEAM["course"], size=13, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0x55, 0x5F, 0x6D))
    para(d, "Project Phase-I", size=13, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0x55, 0x5F, 0x6D))
    d.add_paragraph()
    para(d, title, bold=True, size=22, align=WD_ALIGN_PARAGRAPH.CENTER, color=ACCENT)
    para(d, RC.TEAM["project_title"], size=13, align=WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(4):
        d.add_paragraph()
    para(d, "Submitted by", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for s in RC.TEAM["students"]:
        para(d, f"{s['name']}  —  {s['reg']}", bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    d.add_paragraph()
    para(d, f"Course Instructor: {RC.TEAM['instructor']}", align=WD_ALIGN_PARAGRAPH.CENTER)
    para(d, f"GitHub: {RC.TEAM['repo']}", size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    page_break(d)


def toc(d):
    d.add_heading("Contents", level=1)
    p = d.add_paragraph()
    run = p.add_run()
    for tag, text in (("begin", None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    p.add_run("Right-click here and choose “Update Field” to build the table of contents.").italic = True
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    p.add_run()._r.append(end)
    page_break(d)


# ---------------------------------------------------------------------------
# documents
# ---------------------------------------------------------------------------

def build_report():
    d = new_doc()
    title_page(d, RC.TEAM["short_name"])
    toc(d)
    sec_abstract(d)
    d.add_heading("Introduction", level=1)
    para(d, "Problem statement. Cloud teams declare infrastructure in Terraform, but the live AWS account drifts "
            "away from that declaration through console edits, hot-fixes and automation. Drift in security "
            "settings – public S3 buckets, internet-facing admin ports, over-privileged IAM roles – is a leading "
            "cause of cloud breaches, yet today's tools either never look at the live account or report every "
            "drift with the same urgency and no explanation.", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    para(d, "Existing challenges. Static scanners miss post-deployment changes; terraform plan and AWS Config show "
            "differences without severity; ML and LLM approaches are opaque, costly or hallucinate; automated "
            "responses are rare and unsafe.", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    para(d, "Proposed solution. InfraDriftGuard continuously compares Terraform's desired state with live AWS "
            "configuration, classifies every drift into four risk levels with an explainable Random Forest, and "
            "responds through a serverless AWS pipeline (alerts, storage, optional safe remediation, secured API "
            "and dashboard).", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    sec_literature(d)
    for s in RC.TEAM["students"]:
        sec_research_gap(d, s)
    sec_combined_gap(d)
    page_break(d)
    sec_objectives(d)
    sec_novelty(d)
    page_break(d)
    sec_architecture(d)
    page_break(d)
    sec_dataset(d)
    page_break(d)
    sec_aws(d)
    sec_implementation(d)
    page_break(d)
    sec_results(d)
    page_break(d)
    sec_deploy(d)
    sec_github(d)
    sec_conclusion(d)
    page_break(d)
    sec_references(d)
    d.save(os.path.join(DOCS, "Project_Report.docx"))


def build_single(filename, title, *section_fns, subtitle=None):
    d = new_doc()
    doc_header(d, title, subtitle)
    for fn in section_fns:
        fn(d)
    d.save(filename)


def main():
    os.makedirs(DOCS, exist_ok=True)
    build_report()

    def lit(d):
        d.add_paragraph()
        sec_literature(d, h=2, landscape=False)
        d.add_heading("References", level=2)
        for i, p in enumerate(RC.PAPERS, 1):
            para(d, f"[{i}] {p['ref']}", size=9.5)

    d = new_doc()
    d.sections[0].orientation = WD_ORIENT.LANDSCAPE
    d.sections[0].page_width, d.sections[0].page_height = Cm(29.7), Cm(21)
    d.sections[0].left_margin = d.sections[0].right_margin = Cm(1.6)
    doc_header(d, "Literature Survey")
    lit(d)
    d.save(os.path.join(DOCS, "Literature_Survey.docx"))

    s1, s2 = RC.TEAM["students"]
    build_single(os.path.join(DOCS, "Research_Gap_Student1.docx"), "Research Gap Analysis — Student 1",
                 lambda d: sec_research_gap(d, s1, h=2))
    build_single(os.path.join(DOCS, "Research_Gap_Student2.docx"), "Research Gap Analysis — Student 2",
                 lambda d: sec_research_gap(d, s2, h=2))
    build_single(os.path.join(DOCS, "Research_Gap.docx"), "Research Gap Analysis",
                 lambda d: sec_research_gap(d, s1, h=2), lambda d: sec_research_gap(d, s2, h=2),
                 lambda d: sec_combined_gap(d, h=2))
    build_single(os.path.join(DOCS, "Abstract.docx"), "Abstract", lambda d: para(d, RC.ABSTRACT,
                 align=WD_ALIGN_PARAGRAPH.JUSTIFY), subtitle=f"{len(RC.ABSTRACT.split())} words")
    build_single(os.path.join(DOCS, "Objectives.docx"), "Project Objectives", lambda d: sec_objectives(d, h=2))
    build_single(os.path.join(DOCS, "Novelty.docx"), "Novelty Summary", lambda d: sec_novelty(d, h=2))
    build_single(os.path.join(DOCS, "AWS_Services_Planning.docx"), "AWS Services Planning", lambda d: sec_aws(d, h=2))
    build_single(os.path.join(ROOT, "dataset", "Dataset_Details.docx"), "Dataset Details",
                 lambda d: sec_dataset(d, h=2))
    print("Built:", ", ".join(sorted(os.listdir(DOCS))))


if __name__ == "__main__":
    main()
